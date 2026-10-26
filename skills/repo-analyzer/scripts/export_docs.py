#!/usr/bin/env python3
"""
export_docs.py — Word (.docx) and PDF versions of the repo-analyzer report.

One document model is built from facts.json + data/<lang>/*.json and written twice:
  <dest>/<repo>-analysis-<lang>.docx   python-docx; diagrams as PNG; cover, TOC field, headers/footers, page numbers
  <dest>/<repo>-analysis-<lang>.pdf    headless Chromium; diagrams as vector SVG; cover, TOC with page numbers, bookmarks

Mermaid diagrams are rendered once per language by Mermaid itself inside headless Chromium (Playwright),
so the documents contain finished pictures and open offline.

Requirements:  pip install python-docx playwright pypdf && python -m playwright install chromium
               (pypdf is optional: without it the PDF has no separate cover and no page numbers in the TOC)

Normally called by render_report.py (--formats docx,pdf). Stand-alone:
  python export_docs.py --out DIR [--langs en,tr] [--formats docx,pdf] [--mermaid-js PATH|URL]
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import io
import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import render_report as rr  # noqa: E402

esc = rr.esc
CDN_MERMAID = "https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.min.js"
FONT_LINK = ("https://fonts.googleapis.com/css2?family=Archivo:wdth,wght@62..125,300..800"
             "&family=JetBrains+Mono:wght@400;600&display=swap")
SANS = '"Archivo", "Segoe UI", "Helvetica Neue", Arial, "Liberation Sans", Carlito, "DejaVu Sans", sans-serif'
MONO = '"JetBrains Mono", Consolas, "DejaVu Sans Mono", monospace'
INK, INK2, TEAL, OCHRE, BRICK, RULE = "17232C", "4A5A60", "1D6B67", "A8701A", "9E3A2F", "C9D0C8"


def logo() -> bytes | None:
    """Brand logo PNG from REPO_ANALYZER_LOGO (set by --logo); None when unset/missing."""
    p = os.environ.get("REPO_ANALYZER_LOGO")
    return Path(p).read_bytes() if p and Path(p).is_file() else None


def logo_img(style: str) -> str:
    b = logo()
    return f'<img src="data:image/png;base64,{base64.b64encode(b).decode()}" style="{style}">' if b else ""


SEV_COLOR = {"critical": "7A2A22", "high": BRICK, "medium": OCHRE, "low": "5B7A4E", "info": INK2}


def missing_deps(formats) -> str:
    need = []
    try:
        import playwright.sync_api  # noqa: F401
    except ImportError:
        need.append("playwright")
    if "docx" in formats:
        try:
            import docx  # noqa: F401
        except ImportError:
            need.append("python-docx")
    if not need:
        return ""
    return f"missing {', '.join(need)} → pip install {' '.join(need)} pypdf && python -m playwright install chromium"


def cap(s: str) -> str:
    return s[:1].upper() + s[1:] if s else s


def sev_key(x: dict) -> int:
    s = x.get("severity")
    return rr.SEVERITIES.index(s) if s in rr.SEVERITIES else 9


# ============================================================================ document model
#
# blocks: ("h", level, text, anchor) ("p", runs, style) ("bullets", [runs]) ("table", headers, rows, widths, num_cols)
#         ("figure", key, caption, derived, source) ("code", text)
# runs:   str | ("b", s) | ("i", s) | ("muted", s) | ("code", s) | ("ref", s, ok, url) | ("sev", enum, label)


class Doc:
    def __init__(self, facts: dict, t: dict, lang: str, refs: rr.Refs):
        self.facts, self.t, self.lang, self.refs = facts, t, lang, refs
        self.blocks: list = []
        self.figs: dict[str, str] = {}
        self.n1 = self.n2 = self.nfig = 0
        self.ref_ok, self.ref_bad = 0, []

    # ---- block helpers
    def h(self, level: int, text: str):
        if level == 1:
            self.n1, self.n2 = self.n1 + 1, 0
            text = f"{self.n1}. {text}"
        elif level == 2:
            self.n2 += 1
            text = f"{self.n1}.{self.n2} {text}"
        self.blocks.append(("h", level, text, f"sec{len(self.blocks)}"))

    def p(self, runs, style: str | None = None):
        runs = [r for r in (runs if isinstance(runs, list) else [runs]) if r]
        if runs:
            self.blocks.append(("p", runs, style))

    def bullets(self, items):
        items = [i for i in items if i]
        if items:
            self.blocks.append(("bullets", items))

    def table(self, headers, rows, widths, num_cols=()):
        if rows:
            self.blocks.append(("table", headers, rows, widths, set(num_cols)))

    def code(self, text: str):
        if text:
            self.blocks.append(("code", text))

    def prov(self, derived: bool, source: str = ""):
        self.p([self.t["prov_derived"] if derived else self.t["prov_model"].format(src=source)],
               "prov-derived" if derived else "prov-model")

    def figure(self, code: str, caption: str, derived: bool, source: str = ""):
        code = code if code.lstrip().startswith("<svg") else rr.strip_fences(code)
        if not code:
            return
        key = hashlib.sha1(code.encode()).hexdigest()[:16]
        self.figs[key] = code
        self.nfig += 1
        self.blocks.append(("figure", key, f'{self.t["figure"]} {self.nfig}: {caption}', derived, source))

    def empty(self, file: str):
        self.p([self.t["empty"].format(file=f"data/{self.lang}/{file}")], "note")

    # ---- run helpers
    def ref(self, r: str):
        if not isinstance(r, str) or not r.strip():
            return None
        ok, _ = self.refs.status(r)
        if ok:
            self.ref_ok += 1
        else:
            self.ref_bad.append(r)
        url = self.refs.url(r) if ok else None
        return ("ref", r, ok, url if url and url.startswith("http") else None)

    def refs_runs(self, refs) -> list:
        out = []
        for r in refs or []:
            x = self.ref(r)
            if x:
                out += [", "] if out else []
                out.append(x)
        return out

    def e(self, v) -> str:
        return rr.te(self.t, v)


def build(facts: dict, data: dict, t: dict, lang: str, refs: rr.Refs) -> Doc:
    d = Doc(facts, t, lang, refs)
    ov = data.get("overview.json") or {}
    arch = data.get("architecture.json") or {}
    fs = data.get("features.json") or {}
    flows = (data.get("flows") or {}).get("flows", [])
    hl = data.get("health.json") or {}
    tot, git, infra = facts["totals"], facts.get("git") or {}, facts["infra"]
    findings = sorted([x for x in hl.get("findings", []) if isinstance(x, dict)], key=sev_key)
    feats = [x for x in fs.get("features", []) if isinstance(x, dict)]
    names = {x.get("id"): x.get("name") for x in feats}

    # ------------------------------------------------------------------ 1 executive summary
    d.h(1, t["h_exec"])
    if ov:
        d.prov(False, "orchestrator, quality-auditor")
        d.p([ov.get("one_liner", "")], "lead")
        d.p([ov.get("purpose", "")])
        if ov.get("audience"):
            d.p([("b", f'{t["audience"]}: '), ov["audience"]])
    else:
        d.empty("overview.json")
    if hl.get("summary"):
        d.p([hl["summary"]])

    d.h(2, t["key_figures"])
    langs = [x for x in facts["languages"] if x["lines"] > 0]
    rows = [[cap(t["files"]), f'{tot["files"]:,}'], [cap(t["lines"]), f'{tot["lines"]:,}'],
            [cap(t["languages"]), ", ".join(x["lang"] for x in langs[:4]) + (f" (+{len(langs) - 4})" if len(langs) > 4 else "")]]
    if git.get("available"):
        rows += [[cap(t["commits"]), f'{git.get("total_commits", 0):,}'],
                 [cap(t["contributors"]), str(git.get("authors_in_window", 0))]]
    rows += [[cap(t["tests"]), str(tot["test_ratio"])], [cap(t["routes_n"]), str(len(facts.get("routes", [])))]]
    if feats:
        rows.append([cap(t["features_n"]), f'{len(feats)} ({len(flows)} {t["flows_n"]})'])
    if findings:
        counts = [f'{sum(1 for x in findings if x.get("severity") == s)} {d.e(s).lower()}'
                  for s in rr.SEVERITIES if any(x.get("severity") == s for x in findings)]
        rows.append([cap(t["findings_n"]), ", ".join(counts)])
    d.table([t["metric"], t["value"]], [[[a], [b]] for a, b in rows], [0.38, 0.62])

    top = [x for x in findings if x.get("severity") in ("critical", "high", "medium")][:6]
    if top:
        d.h(2, t["top_findings"])
        d.table([t["c_severity"], t["c_finding"], t["c_category"]],
                [[[("sev", x.get("severity"), d.e(x.get("severity")))], [x.get("title", "")], [d.e(x.get("category"))]]
                 for x in top], [0.16, 0.62, 0.22])
    if hl.get("strengths"):
        d.h(2, t["strengths"])
        d.bullets([[("b", x.get("title", "")), (" " + x.get("detail", "")) if x.get("detail") else ""]
                   for x in hl["strengths"]])
    if git.get("available") and git.get("monthly_commits"):
        d.h(2, t["activity"])
        d.figure(rr.activity_svg(git["monthly_commits"]), t["activity"], True)

    # ------------------------------------------------------------------ 2 technology
    d.h(1, t["h_tech"])
    if ov.get("tech_stack"):
        d.h(2, t["stack"])
        d.prov(False, "orchestrator")
        d.table([t["c_name"], t["c_category"], t["c_role"], t["c_refs"]],
                [[[("b", s.get("name", ""))], [d.e(s.get("category", ""))], [s.get("role", "")], d.refs_runs(s.get("refs"))]
                 for s in ov["tech_stack"]], [0.24, 0.18, 0.26, 0.32])
    if facts["frameworks"]:
        d.p([("muted", f'{t["derived"]}: '), ", ".join(f'{x["name"]} ({x["category"]})' for x in facts["frameworks"])], "note")
    d.h(2, t["languages"])
    d.prov(True)
    total = sum(x["lines"] for x in langs) or 1
    d.table([t["c_lang"], cap(t["files"]), cap(t["lines"]), t["c_share"]],
            [[[d.e(x["lang"])], [f'{x["files"]:,}'], [f'{x["lines"]:,}'], [f'{100 * x["lines"] / total:.1f}%']] for x in langs[:12]],
            [0.4, 0.2, 0.2, 0.2], num_cols=(1, 2, 3))
    items = [[f"Docker: ", d.ref(x["path"]), ("muted", f' FROM {", ".join(x["from"])}')] for x in infra["dockerfiles"][:6]]
    items += [["Compose ", d.ref(x["path"]), f': {", ".join(x["services"])}'] for x in infra["compose_services"][:3]]
    items += [["CI ", d.ref(x["path"]), ("muted", f' {x["system"]}')] for x in infra["ci"][:8]]
    if infra["k8s_kinds"]:
        items.append([f'Kubernetes: {", ".join(f"{k} ×{v}" for k, v in infra["k8s_kinds"].items())}'])
    if infra["terraform_resources"]:
        items.append([f'Terraform: {", ".join(f"{k} ×{v}" for k, v in list(infra["terraform_resources"].items())[:12])}'])
    if items:
        d.h(2, t["infra"])
        d.prov(True)
        d.bullets(items)
    if infra["env_keys"]:
        d.h(2, t["envkeys"])
        for e in infra["env_keys"][:3]:
            d.p([d.ref(e["path"])])
            d.code(", ".join(e["keys"]))

    # ------------------------------------------------------------------ 3 architecture
    d.h(1, t["h_arch"])
    if arch:
        d.prov(False, "repo-architect")
        d.p([arch.get("summary", "")])
        if arch.get("style"):
            d.p([("b", f'{t["style"]}: '), arch["style"]])
        dg = arch.get("diagrams", {})
        for key in ("context", "containers", "components"):
            if dg.get(key):
                d.h(2, t[key])
                d.figure(dg[key], t[key], False, "repo-architect")
        if arch.get("layers"):
            d.h(2, t["layers"])
            d.table([t["c_name"], t["c_resp"], t["c_refs"]],
                    [[[("b", x.get("name", ""))], [x.get("responsibility", "")], d.refs_runs(x.get("refs"))]
                     for x in arch["layers"]], [0.22, 0.48, 0.30])
        if arch.get("external_systems"):
            d.h(2, t["external"])
            d.table([t["c_name"], t["c_kind"], t["c_purpose"], t["c_refs"]],
                    [[[("b", x.get("name", ""))], [d.e(x.get("kind", ""))], [x.get("purpose", "")], d.refs_runs(x.get("refs"))]
                     for x in arch["external_systems"]], [0.2, 0.16, 0.36, 0.28])
    else:
        d.empty("architecture.json")
    graph = rr.dep_graph_mermaid(facts)
    if graph:
        d.h(2, t["depgraph"])
        d.prov(True)
        d.figure(graph, t["depgraph"], True)
    if arch:
        dm = arch.get("data_model") or {}
        if dm.get("diagram") or dm.get("entities"):
            d.h(2, t["datamodel"])
            if dm.get("diagram"):
                d.figure(dm["diagram"], t["datamodel"], False, "repo-architect")
            d.table([t["c_entity"], t["c_desc"], t["c_refs"]],
                    [[[("b", x.get("name", ""))], [x.get("description", "")], d.refs_runs(x.get("refs"))]
                     for x in dm.get("entities", [])], [0.2, 0.5, 0.3])
        if arch.get("key_decisions"):
            d.h(2, t["decisions"])
            for x in arch["key_decisions"]:
                d.blocks.append(("h", 3, x.get("decision", ""), f"sec{len(d.blocks)}"))
                if x.get("rationale"):
                    d.p([("b", f'{t["rationale"]}: '), x["rationale"]])
                if x.get("tradeoffs"):
                    d.p([("b", f'{t["tradeoffs"]}: '), x["tradeoffs"]])
                d.p(d.refs_runs(x.get("refs")), "refs")

    # ------------------------------------------------------------------ 4 features
    d.h(1, t["h_features"])
    if fs:
        d.prov(False, "feature-mapper")
        if fs.get("feature_map"):
            d.h(2, t["featuremap"])
            d.figure(fs["feature_map"], t["featuremap"], False, "feature-mapper")
        d.h(2, t["h_features"])
        d.table([t["c_feature"], t["c_category"], t["c_conf"], t["c_summary"]],
                [[[("b", x.get("name", ""))], [d.e(x.get("category", ""))], [d.e(x.get("confidence", ""))], [x.get("summary", "")]]
                 for x in feats], [0.22, 0.17, 0.11, 0.50])
        for x in feats:
            d.blocks.append(("h", 3, x.get("name", ""), f"sec{len(d.blocks)}"))
            d.p([("muted", f'{d.e(x.get("category", ""))} · {t["confidence"]}: {d.e(x.get("confidence", "?"))}')], "note")
            d.p([x.get("summary", "")])
            if x.get("entry_points"):
                d.table([t["c_refs"], t["c_kind"], t["c_detail"]],
                        [[[d.ref(e.get("ref", ""))], [d.e(e.get("kind", ""))], [("code", e.get("detail", ""))]]
                         for e in x["entry_points"]], [0.36, 0.2, 0.44])
            if x.get("key_files"):
                d.p([("b", f'{t["keyfiles"]}: ')] + d.refs_runs(x["key_files"]))
            if x.get("depends_on"):
                d.p([("b", f'{t["dependson"]}: '), ", ".join(names.get(i, i) for i in x["depends_on"])])
    else:
        d.empty("features.json")

    # ------------------------------------------------------------------ 5 flows
    d.h(1, t["h_flows"])
    if not flows:
        d.empty("flows-*.json")
    else:
        d.prov(False, "flow-tracer")
    for fl in flows:
        d.h(2, fl.get("name", ""))
        if fl.get("feature_id"):
            d.p([("muted", f'{t["flow_for"]}: {names.get(fl["feature_id"], fl["feature_id"])}')], "note")
        if fl.get("trigger"):
            d.p([("b", f'{t["trigger"]}: '), fl["trigger"]])
        if fl.get("sequence"):
            d.figure(fl["sequence"], fl.get("name", ""), False, "flow-tracer")
        if fl.get("steps"):
            d.blocks.append(("h", 3, t["steps"], f"sec{len(d.blocks)}"))
            d.table(["#", t["c_actor"], t["c_action"], t["c_refs"]],
                    [[[str(i)], [("b", s.get("actor", ""))], [s.get("action", "")], [d.ref(s.get("ref", ""))]]
                     for i, s in enumerate(fl["steps"], 1)], [0.06, 0.2, 0.46, 0.28], num_cols=(0,))
        if fl.get("error_paths"):
            d.blocks.append(("h", 3, t["errors"], f"sec{len(d.blocks)}"))
            d.table([t["c_condition"], t["c_handling"], t["c_refs"]],
                    [[[e.get("condition", "")], [e.get("handling", "")], [d.ref(e.get("ref", ""))]] for e in fl["error_paths"]],
                    [0.3, 0.42, 0.28])
        if fl.get("side_effects"):
            d.blocks.append(("h", 3, t["sidefx"], f"sec{len(d.blocks)}"))
            d.bullets([[s] for s in fl["side_effects"]])
        if fl.get("notes"):
            d.blocks.append(("h", 3, t["notes"], f"sec{len(d.blocks)}"))
            d.p([fl["notes"]])

    # ------------------------------------------------------------------ 6 health
    d.h(1, t["h_health"])
    if hl:
        d.prov(False, "quality-auditor")
        d.p([hl.get("summary", "")])
        if findings:
            d.h(2, t["findings"])
            d.table([t["c_severity"], t["c_finding"], t["c_category"]],
                    [[[("sev", x.get("severity"), d.e(x.get("severity")))], [x.get("title", "")], [d.e(x.get("category"))]]
                     for x in findings], [0.16, 0.62, 0.22])
            for x in findings:
                d.blocks.append(("h", 3, x.get("title", ""), f"sec{len(d.blocks)}"))
                d.p([("sev", x.get("severity"), d.e(x.get("severity"))), ("muted", f'  {d.e(x.get("category"))}')])
                d.p([x.get("detail", "")])
                if x.get("recommendation"):
                    d.p([("b", f'{t["recommendation"]}: '), x["recommendation"]])
                d.p(d.refs_runs(x.get("refs")), "refs")
    else:
        d.empty("health.json")
    tst = hl.get("testing") or {}
    d.h(2, t["testing"])
    d.p([f'{tot["test_files"]} test {t["files"]}, {tot["test_lines"]:,} / {tot["source_lines"]:,} {t["lines"]} '
         f'({t["tests"]}: {tot["test_ratio"]})'])
    if tst.get("summary"):
        d.p([tst["summary"]])
    if tst.get("frameworks"):
        d.p([("b", f'{t["frameworks"]}: '), ", ".join(tst["frameworks"])])
    if tst.get("gaps"):
        d.p([("b", t["gaps"])])
        d.bullets([[g.get("area", ""), " "] + d.refs_runs(g.get("refs")) for g in tst["gaps"]])
    if facts.get("hotspots"):
        d.h(2, t["hotspots"])
        d.prov(True)
        notes = {n.get("ref"): n.get("why_it_matters") for n in hl.get("hotspot_notes", [])}
        d.table([t["c_file"], cap(t["commits"]), cap(t["lines"]), t["fixes"], cap(t["authors"]), t["score"]],
                [[[d.ref(x["path"])] + ([("muted", " " + notes[x["path"]])] if notes.get(x["path"]) else []),
                  [str(x["commits"])], [str(x["lines"])], [str(x["fix_commits"])], [str(x["authors"])], [f'{x["score"]:.2f}']]
                 for x in facts["hotspots"][:15]], [0.45, 0.11, 0.11, 0.11, 0.11, 0.11], num_cols=(1, 2, 3, 4, 5))
    if git.get("coupling"):
        d.h(2, t["coupling"])
        d.prov(True)
        d.table([t["c_file"], t["c_file"], t["shared"], t["degree"]],
                [[[d.ref(c["a"])], [d.ref(c["b"])], [str(c["shared_commits"])], [f'{c["degree"]}']] for c in git["coupling"][:12]],
                [0.36, 0.36, 0.14, 0.14], num_cols=(2, 3))
    if git.get("bus_factor"):
        d.h(2, t["busfactor"])
        d.prov(True)
        d.table([cap(t["dir"]), t["bf"], t["topauthor"], t["share"], cap(t["authors"])],
                [[[("code", b["dir"])], [str(b["bus_factor"])], [b["top_author"]], [f'{b["top_share"]:.0%}'], [str(b["authors"])]]
                 for b in git["bus_factor"][:12]], [0.34, 0.14, 0.26, 0.12, 0.14], num_cols=(1, 3, 4))
    elif not git.get("available"):
        d.p([t["nogit"]], "note")
    mk = facts["markers"]
    if mk["total"]:
        d.h(2, f'{t["markers"]} ({mk["total"]})')
        d.table([t["c_refs"], t["c_marker"]], [[[d.ref(x["ref"])], [("code", x["text"])]] for x in mk["samples"][:15]], [0.35, 0.65])
    if facts.get("secrets_candidates"):
        d.h(2, t["secrets"])
        d.table([t["c_refs"], t["c_kind"]], [[[d.ref(s["ref"])], [s["kind"]]] for s in facts["secrets_candidates"][:20]], [0.5, 0.5])
    if facts.get("largest_files"):
        d.h(2, t["largest"])
        d.table([t["c_file"], cap(t["lines"])], [[[d.ref(x["path"])], [f'{x["lines"]:,}']] for x in facts["largest_files"][:10]],
                [0.8, 0.2], num_cols=(1,))

    # ------------------------------------------------------------------ 7 onboarding
    d.h(1, t["h_onboarding"])
    if ov:
        d.prov(False, "orchestrator")
        if ov.get("how_to_run"):
            d.h(2, t["howtorun"])
            d.table(["#", t["c_step"], t["c_command"], t["c_refs"]],
                    [[[str(i)], [s.get("step", "")], [("code", s.get("command", ""))], [d.ref(s.get("ref", ""))]]
                     for i, s in enumerate(ov["how_to_run"], 1)], [0.06, 0.3, 0.4, 0.24], num_cols=(0,))
        if ov.get("reading_order"):
            d.h(2, t["reading"])
            d.table(["#", t["c_file"], t["c_why"]],
                    [[[str(i)], [d.ref(r.get("ref", ""))], [r.get("why", "")]] for i, r in enumerate(ov["reading_order"], 1)],
                    [0.06, 0.38, 0.56], num_cols=(0,))
        if ov.get("glossary"):
            d.h(2, t["glossary"])
            d.table([t["c_term"], t["c_meaning"], t["c_refs"]],
                    [[[("b", g.get("term", ""))], [g.get("meaning", "")], d.refs_runs(g.get("refs"))] for g in ov["glossary"]],
                    [0.22, 0.5, 0.28])
        if ov.get("starter_tasks"):
            d.h(2, t["tasks"])
            d.p([t["tasks_note"]], "note")
            for x in ov["starter_tasks"]:
                d.blocks.append(("h", 3, x.get("title", ""), f"sec{len(d.blocks)}"))
                d.p([("muted", f'{t["c_type"]}: {d.e(x.get("type", ""))} · {t["c_diff"]}: {d.e(x.get("difficulty", ""))}')], "note")
                d.p([x.get("description", "")])
                d.p(d.refs_runs(x.get("refs")), "refs")
                if x.get("acceptance"):
                    d.p([("b", t["acceptance"])])
                    d.bullets([[c] for c in x["acceptance"]])
    else:
        d.empty("overview.json")

    # ------------------------------------------------------------------ 8 appendices
    cf = data.get("conformance.json") or {}
    if cf:
        d.h(1, t["h_conformance"])
        doc = cf.get("document") or {}
        d.p([("b", doc.get("title", "")), f'  {doc.get("source", "")}'])
        d.p([("muted", t["c_notsot"])])
        d.p([cf.get("summary", "")])
        order = {v: i for i, v in enumerate(rr.VERDICTS)}
        claims = sorted(cf.get("claims", []), key=lambda x: order.get(x.get("verdict"), 9))
        if claims:
            d.h(2, t["c_claims"])
            d.table([t["c_verdict"], t["c_claims"], t["c_docref"]],
                    [[[t.get("v_" + str(x.get("verdict")), "")], [x.get("claim", "")], [d.e(x.get("doc_ref", ""))]]
                     for x in claims], [0.18, 0.62, 0.20])
            for x in claims:
                d.blocks.append(("h", 3, x.get("claim", ""), f"sec{len(d.blocks)}"))
                d.p([("b", t.get("v_" + str(x.get("verdict")), "")),
                     ("muted", f'  {d.e(x.get("doc_ref", ""))}')])
                d.p([("b", f'{t["c_finding"]}: '), x.get("finding", "")])
                d.p(d.refs_runs(x.get("refs")), "refs")
        if cf.get("gaps"):
            d.h(2, t["c_gaps"])
            for g in cf["gaps"]:
                d.p([("b", g.get("title", "") + ": "), g.get("detail", "")])
                d.p(d.refs_runs(g.get("refs")), "refs")

    fs = data.get("feasibility.json") or {}
    if fs:
        d.h(1, t["h_feasibility"])
        d.p([fs.get("summary", "")])
        if fs.get("verdict"):
            d.p([("b", fs["verdict"])])
        order = {v: i for i, v in enumerate(rr.RATINGS)}
        items = sorted(fs.get("expectations", []), key=lambda x: order.get(x.get("rating"), 9))
        if items:
            d.h(2, t["f_expect"])
            d.table([t["f_rating"], t["f_expect"], t["f_effort"]],
                    [[[t.get("r_" + str(x.get("rating")), "")], [x.get("expectation", "")], [d.e(x.get("effort", ""))]]
                     for x in items], [0.22, 0.58, 0.20])
            for x in items:
                d.blocks.append(("h", 3, x.get("expectation", ""), f"sec{len(d.blocks)}"))
                d.p([("b", t.get("r_" + str(x.get("rating")), "")),
                     ("muted", f'  {d.e(x.get("doc_ref", ""))}')])
                d.p([x.get("assessment", "")])
                if x.get("prerequisites"):
                    d.p([("b", t["f_prereq"] + ":")])
                    d.bullets([[p] for p in x["prerequisites"]])
                if x.get("risk"):
                    d.p([("b", t["f_risk"] + ": "), x["risk"]])
                d.p(d.refs_runs(x.get("evidence_refs")), "refs")
        if fs.get("critical_path"):
            d.h(2, t["f_path"])
            d.bullets([[("b", c.get("step", "") + " "), c.get("why", "")] for c in fs["critical_path"]])
        if fs.get("open_questions"):
            d.h(2, t["f_open"])
            d.bullets([[("b", q.get("question", "") + " "), q.get("why_it_matters", "")] for q in fs["open_questions"]])

    rd = data.get("readiness.json") or {}
    if rd:
        d.h(1, t["h_readiness"])
        d.p([rd.get("summary", "")])
        if rd.get("verdict"):
            d.p([("b", rd["verdict"])])
        for fig in rd.get("diagrams", []):
            if fig.get("diagram"):
                d.figure(fig["diagram"], fig.get("title", ""), False, "assessment")
        if rd.get("mvp_scope"):
            d.h(2, t["rd_scope"])
            d.table([t["c_detail"], t["c_kind"], t["c_why"]],
                    [[[i.get("item", "")], [t.get("rd_" + str(i.get("decision")), "")], [i.get("why", "")]]
                     for i in rd["mvp_scope"]], [0.38, 0.14, 0.48])
        for dec in rd.get("decisions", []):
            d.blocks.append(("h", 2, dec.get("question", ""), f"sec{len(d.blocks)}"))
            if dec.get("impact"):
                d.p([("muted", f'{t["rd_impact"]}: {d.e(dec.get("impact", ""))}')])
            if dec.get("background"):
                d.p([dec["background"]])
            if dec.get("diagram"):
                d.figure(dec["diagram"], dec.get("question", ""), False, "assessment")
            if dec.get("options"):
                d.p([("b", t["rd_options"] + ":")])
                d.bullets([[("b", o.get("option", "") + " "), o.get("tradeoffs", "")] for o in dec["options"]])
            if dec.get("recommendation"):
                d.p([("b", t["rd_reco"] + ": "), dec["recommendation"]])
            d.p(d.refs_runs(dec.get("refs")), "refs")
        for w in rd.get("workstreams", []):
            d.blocks.append(("h", 2, w.get("title", ""), f"sec{len(d.blocks)}"))
            d.p([("muted", f'{t["rd_outcome"]}: {d.e(w.get("outcome", ""))}')])
            if w.get("tasks"):
                d.bullets([[task] for task in w["tasks"]])
            d.p(d.refs_runs(w.get("refs")), "refs")
        if rd.get("sequence"):
            d.h(2, t["rd_seq"])
            d.bullets([[("b", s.get("step", "") + " "), s.get("why", "")] for s in rd["sequence"]])
        if rd.get("open_questions"):
            d.h(2, t["rd_open"])
            d.bullets([[("b", q.get("question", "") + " "), q.get("why_it_matters", "")] for q in rd["open_questions"]])

    d.h(1, t["h_appendix"])
    routes = facts.get("routes", [])
    if routes:
        d.h(2, f'{t["app_routes"]} ({len(routes)})')
        d.prov(True)
        d.table([t["c_method"], t["c_path"], t["c_refs"], t["c_framework"]],
                [[[("code", r["method"])], [("code", r["path"])], [d.ref(r["ref"])], [r["framework"]]] for r in routes[:300]],
                [0.12, 0.38, 0.32, 0.18])
    d.h(2, t["app_method"])
    for para in t["method_p"]:
        d.p([para])
    bad = sorted(set(d.ref_bad))
    d.p([("b", t["verified"].format(ok=d.ref_ok, bad=len(d.ref_bad)))])
    if bad:
        d.p([("b", t["missing_list"])])
        d.bullets([[("ref", r, False, None)] for r in bad[:40]])
    return d


def meta_of(facts: dict, data: dict, t: dict, lang: str) -> dict:
    git = facts.get("git") or {}
    ov = data.get("overview.json") or {}
    rows = [(t["repository"], git.get("remote_web") or facts["repo"]["name"])]
    if git.get("available"):
        rows += [("Commit", (git.get("head") or "")[:12]), (t["branch"], git.get("branch", ""))]
    rows += [(t["generated"], facts["generated_at"][:10]), (t["report_lang"], t["lang_name"]),
             (cap(t["files"]), f'{facts["totals"]["files"]:,} · {facts["totals"]["lines"]:,} {t["lines"]}')]
    return {"name": facts["repo"]["name"], "title": t["report_title"], "kicker": t["report_kicker"],
            "one_liner": ov.get("one_liner", ""), "date": facts["generated_at"][:10], "rows": rows, "lang": lang}


# ============================================================================ diagram rendering (Chromium)

MERMAID_CFG = {
    "startOnLoad": False, "securityLevel": "strict", "theme": "base",
    "flowchart": {"curve": "basis", "htmlLabels": True, "padding": 12},
    "sequence": {"mirrorActors": False, "messageAlign": "left", "actorMargin": 30, "width": 130, "diagramMarginX": 8,
                 "diagramMarginY": 8, "boxMargin": 6, "noteMargin": 8},
    "themeVariables": {
        "fontFamily": SANS.replace('"', "'"), "fontSize": "14px",
        "primaryColor": "#FBFCFA", "primaryTextColor": "#17232C", "primaryBorderColor": "#17232C", "lineColor": "#4A5A60",
        "secondaryColor": "#E3EAE2", "tertiaryColor": "#F3F5F0", "clusterBkg": "#F3F5F0", "clusterBorder": "#8FA08F",
        "edgeLabelBackground": "#FFFFFF", "actorBkg": "#FBFCFA", "actorBorder": "#17232C", "actorTextColor": "#17232C",
        "signalColor": "#17232C", "signalTextColor": "#17232C", "noteBkgColor": "#F3E6C8", "noteBorderColor": "#A8701A",
        "labelBoxBkgColor": "#FBFCFA", "activationBkgColor": "#E3EAE2"},
}
SVG_CSS = (f".activity rect{{fill:#{TEAL}}}.activity line{{stroke:#{INK}}}.activity text{{font:11px {SANS};fill:#{INK2}}}"
           ".contours path{fill:none;stroke:#B7C2B8;stroke-width:1;stroke-linejoin:round}"
           ".contours path.idx{stroke:#93A294;stroke-width:1.8}.contours path.peak{fill:#A8701A;stroke:none}")
RENDER_PAGE = f"""<!doctype html><html><head><meta charset="utf-8"><link href="{FONT_LINK}" rel="stylesheet">
<style>html,body{{margin:0;background:#fff;font-family:{SANS}}}#stage{{display:inline-block;padding:14px;background:#fff}}
#stage svg{{display:block}}{SVG_CSS}</style></head><body><div id="stage"></div></body></html>"""
PLACE_JS = """(svg) => {
  const st = document.getElementById('stage'); st.innerHTML = svg;
  const el = st.querySelector('svg'); const vb = el.viewBox && el.viewBox.baseVal;
  let w = vb && vb.width ? vb.width : el.getBoundingClientRect().width;
  let h = vb && vb.height ? vb.height : el.getBoundingClientRect().height;
  if (!(vb && vb.width)) el.setAttribute('viewBox', `0 0 ${w} ${h}`);
  el.setAttribute('width', w); el.setAttribute('height', h); el.style.maxWidth = 'none';
  el.removeAttribute('preserveAspectRatio');
  if (el.classList.contains('contours')) el.setAttribute('preserveAspectRatio', 'xMidYMid slice');
  return {w, h, svg: el.outerHTML};
}"""


def render_assets(browser, figs: dict[str, str], mermaid_js: str | None, want_png: bool) -> dict:
    page = browser.new_page(device_scale_factor=2, viewport={"width": 1400, "height": 1000})
    page.set_content(RENDER_PAGE, wait_until="domcontentloaded")
    src = mermaid_js or CDN_MERMAID
    try:
        page.add_script_tag(url=src) if src.startswith(("http://", "https://")) else page.add_script_tag(path=src)
    except Exception as e:
        raise RuntimeError(f"could not load Mermaid from {src} ({str(e).splitlines()[0]}); "
                           "for offline use pass --mermaid-js /path/to/mermaid.min.js") from None
    page.evaluate("""async () => { await Promise.race([document.fonts.load('14px Archivo'),
                                                        new Promise(r => setTimeout(r, 3000))]); }""")
    page.evaluate("cfg => mermaid.initialize(cfg)", MERMAID_CFG)
    assets = {}
    for i, (key, code) in enumerate(figs.items()):
        try:
            if code.lstrip().startswith("<svg"):
                svg = code
            else:
                svg = page.evaluate("""async ([id, code]) => {
                    try { return (await mermaid.render(id, code)).svg; }
                    finally { document.getElementById('d' + id)?.remove(); } }""", [f"fig{i}", code])
            box = page.evaluate(PLACE_JS, svg)
            png = None
            if want_png:
                page.set_viewport_size({"width": int(max(1400, box["w"] + 60)), "height": int(max(1000, box["h"] + 60))})
                png = page.locator("#stage").screenshot(type="png")
            assets[key] = {"svg": box["svg"], "w": box["w"], "h": box["h"], "png": png}
        except Exception as e:  # Mermaid syntax error: keep the source so the document shows what failed
            msg = str(e).split("\n")[0].replace("Error: ", "")[:300]
            assets[key] = {"error": msg, "code": code}
    page.close()
    return assets


# ============================================================================ PDF (HTML → Chromium print)

PDF_CSS = f"""
@page {{ size: A4; }}
html {{ font: 9.6pt/1.5 {SANS}; color: #{INK}; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
body {{ margin: 0; }}
h1, h2, h3 {{ position: relative; break-after: avoid; }}
h1 {{ font-size: 19pt; font-weight: 800; font-stretch: 112%; line-height: 1.15; margin: 0 0 10pt;
      padding-bottom: 6pt; border-bottom: 1.5pt solid #{INK}; break-before: page; }}
h2 {{ font-size: 13pt; font-weight: 700; color: #{TEAL}; margin: 16pt 0 6pt; }}
h3 {{ font-size: 10.5pt; font-weight: 700; margin: 12pt 0 3pt; }}
p {{ margin: 0 0 5.5pt; orphans: 3; widows: 3; }}
p.lead {{ font-size: 12pt; font-stretch: 104%; line-height: 1.4; }}
p.note, .muted {{ color: #{INK2}; font-size: 8.8pt; }}
p.prov-derived, p.prov-model {{ color: #{INK2}; font-size: 8pt; font-style: italic; padding-left: 6pt;
      border-left: 2pt solid #{RULE}; margin-bottom: 8pt; }}
p.prov-model {{ border-left-style: dashed; border-left-color: #{INK2}; }}
p.refs {{ margin-top: -2pt; }}
ul {{ margin: 0 0 7pt; padding-left: 13pt; }} li {{ margin: 0 0 2.5pt; }}
table {{ width: 100%; border-collapse: collapse; table-layout: fixed; font-size: 8.4pt; line-height: 1.35;
         margin: 3pt 0 10pt; border-top: 1pt solid #{INK}; }}
thead th {{ text-align: left; background: #E8ECE6; font-weight: 700; padding: 3.5pt 5pt; border-bottom: .8pt solid #{INK}; }}
td {{ padding: 3.5pt 5pt; border-bottom: .5pt solid #{RULE}; vertical-align: top; overflow-wrap: anywhere; }}
tr {{ break-inside: avoid; }}
td.num, th.num {{ text-align: right; font-variant-numeric: tabular-nums; }}
.ref {{ font-family: {MONO}; font-size: 7.4pt; color: #{TEAL}; text-decoration: none; overflow-wrap: anywhere; }}
.ref.bad {{ color: #{BRICK}; text-decoration: line-through; }}
code {{ font-family: {MONO}; font-size: 7.8pt; overflow-wrap: anywhere; }}
pre {{ font-family: {MONO}; font-size: 7.8pt; background: #F3F5F0; border: .5pt solid #{RULE}; padding: 5pt 7pt;
       white-space: pre-wrap; overflow-wrap: anywhere; margin: 0 0 8pt; }}
figure {{ margin: 6pt 0 12pt; break-inside: avoid; border: 1pt solid #{INK}; padding: 8pt 8pt 5pt; }}
figure.model {{ border-style: dashed; border-color: #{INK2}; }}
.svgbox svg {{ display: block; margin: 0 auto; max-width: 100% !important; width: auto !important; height: auto !important; max-height: 205mm; }}
figcaption {{ font-size: 8pt; color: #{INK2}; margin-top: 6pt; border-top: .5pt solid #{RULE}; padding-top: 4pt; }}
figcaption b {{ color: #{INK}; }}
.fig-err {{ color: #{BRICK}; font-size: 8.5pt; border-left: 2pt solid #{BRICK}; padding-left: 6pt; margin-bottom: 4pt; }}
.sev {{ display: inline-block; font-weight: 700; font-size: 7pt; text-transform: uppercase; letter-spacing: .05em;
        padding: .5pt 5pt; border-radius: 2pt; color: #fff; }}
{"".join(f".sev-{k}{{background:#{v}}}" for k, v in SEV_COLOR.items())}
.mk {{ position: absolute; left: 0; top: 0; font-size: 1px; line-height: 1px; color: #fff; }}
nav.toc h1 {{ break-before: auto; }}
nav.toc ol {{ list-style: none; margin: 0; padding: 0; }}
nav.toc li a {{ display: flex; align-items: baseline; color: #{INK}; text-decoration: none; }}
nav.toc .dots {{ flex: 1; border-bottom: .8pt dotted #93A294; margin: 0 5pt; transform: translateY(-3pt); }}
nav.toc .pg {{ font-variant-numeric: tabular-nums; min-width: 16pt; text-align: right; }}
nav.toc li.l1 {{ font-weight: 700; margin-top: 7pt; font-size: 10.5pt; }}
nav.toc li.l2 {{ padding-left: 16pt; font-size: 9.2pt; margin-top: 2pt; }}
{SVG_CSS}
.cover {{ position: relative; height: 297mm; overflow: hidden; break-after: page; }}
.cover .contours {{ position: absolute; left: 0; top: 0; width: 210mm; height: 118mm; }}
.cover .inner {{ position: absolute; left: 20mm; right: 20mm; top: 128mm; }}
.cover .kicker {{ color: #{TEAL}; font-weight: 700; letter-spacing: .18em; font-size: 9.5pt; margin: 0 0 8pt; }}
.cover h1.title {{ font-size: 40pt; line-height: 1.02; border: 0; margin: 0 0 10pt; padding: 0; break-before: auto;
                   font-stretch: 118%; overflow-wrap: anywhere; }}
.cover .sub {{ font-size: 13pt; line-height: 1.4; margin: 0 0 16pt; max-width: 150mm; }}
.cover table {{ border-top: 1.5pt solid #{INK}; font-size: 9pt; width: 130mm; table-layout: auto; }}
.cover td:first-child {{ color: #{INK2}; width: 38mm; }}
.cover .legend {{ position: absolute; left: 20mm; right: 20mm; bottom: 16mm; font-size: 8pt; color: #{INK2}; }}
.cover .legend p {{ margin: 0 0 3pt; }}
.cover.inbody {{ height: 255mm; }} .cover.inbody .contours {{ width: 100%; height: 95mm; }}
.cover.inbody .inner {{ left: 0; right: 0; top: 105mm; }} .cover.inbody .legend {{ left: 0; right: 0; bottom: 0; }}
.cover:not(.inbody) {{ break-after: auto; }}
.sw {{ display: inline-block; width: 9pt; height: 9pt; margin-right: 5pt; vertical-align: -1pt; border: 1pt solid #{INK}; }}
.sw.dashed {{ border-style: dashed; border-color: #{INK2}; }}
"""


def html_runs(runs) -> str:
    out = []
    for r in runs:
        if not r:
            continue
        if isinstance(r, str):
            out.append(esc(r))
            continue
        k = r[0]
        if k == "b":
            out.append(f"<b>{esc(r[1])}</b>")
        elif k == "i":
            out.append(f"<i>{esc(r[1])}</i>")
        elif k == "muted":
            out.append(f'<span class="muted">{esc(r[1])}</span>')
        elif k == "code":
            out.append(f"<code>{esc(r[1])}</code>")
        elif k == "sev":
            out.append(f'<span class="sev sev-{esc(r[1])}">{esc(r[2])}</span>')
        elif k == "ref":
            _, text, ok, url = r
            if not ok:
                out.append(f'<span class="ref bad">{esc(text)}</span>')
            elif url:
                out.append(f'<a class="ref" href="{esc(url)}">{esc(text)}</a>')
            else:
                out.append(f'<span class="ref">{esc(text)}</span>')
    return "".join(out)


def cover_html(meta: dict, t: dict, inbody: bool = False) -> str:
    rows = "".join(f"<tr><td>{esc(a)}</td><td>{esc(b)}</td></tr>" for a, b in meta["rows"])
    return (f'<section class="cover{" inbody" if inbody else ""}">{rr.contour_svg(meta["name"], 1200, 675)}<div class="inner">'
            f'{logo_img("height:12mm;display:block;margin:0 0 8mm")}<p class="kicker">{esc(meta["kicker"])}</p><h1 class="title">{esc(meta["name"])}</h1>'
            f'<p class="sub">{esc(meta["one_liner"])}</p><table><tbody>{rows}</tbody></table></div>'
            f'<div class="legend"><p><span class="sw"></span>{esc(t["legend_solid"])}</p>'
            f'<p><span class="sw dashed"></span>{esc(t["legend_dashed"])}</p></div></section>')


def body_html(doc: Doc, assets: dict, meta: dict, pages: dict | None, markers: bool, with_cover: bool) -> str:
    t = doc.t
    toc = "".join(
        f'<li class="l{b[1]}"><a href="#{b[3]}"><span class="t">{esc(b[2])}</span><span class="dots"></span>'
        f'<span class="pg">{(pages or {}).get(b[3], "000") if pages is not None else ""}</span></a></li>'
        for b in doc.blocks if b[0] == "h" and b[1] <= 2)
    parts = [cover_html(meta, t, inbody=True) if with_cover else "",
             f'<nav class="toc"><h1>{esc(t["toc"])}</h1><ol>{toc}</ol></nav>']
    for i, b in enumerate(doc.blocks):
        k = b[0]
        if k == "h":
            mk = f'<span class="mk">zzmk{b[3]}zz</span>' if markers else ""
            parts.append(f'<h{b[1]} id="{b[3]}">{mk}{esc(b[2])}</h{b[1]}>')
        elif k == "p":
            parts.append(f'<p class="{b[2] or ""}">{html_runs(b[1])}</p>')
        elif k == "bullets":
            parts.append("<ul>" + "".join(f"<li>{html_runs(x)}</li>" for x in b[1]) + "</ul>")
        elif k == "code":
            parts.append(f"<pre>{esc(b[1])}</pre>")
        elif k == "table":
            _, headers, rows, widths, nums = b
            cols = "".join(f'<col style="width:{w * 100:.1f}%">' for w in widths)
            head = ("<thead><tr>" + "".join(f'<th{" class=num" if j in nums else ""}>{esc(h)}</th>'
                                             for j, h in enumerate(headers)) + "</tr></thead>") if headers else ""
            body = "".join("<tr>" + "".join(f'<td{" class=num" if j in nums else ""}>{html_runs(c)}</td>'
                                            for j, c in enumerate(row)) + "</tr>" for row in rows)
            parts.append(f"<table><colgroup>{cols}</colgroup>{head}<tbody>{body}</tbody></table>")
        elif k == "figure":
            _, key, caption, derived, source = b
            a = assets.get(key, {})
            prov = t["derived"] if derived else f'{t["model"]}{" (" + source + ")" if source else ""}'
            inner = (f'<div class="svgbox">{a["svg"]}</div>' if "svg" in a else
                     f'<div class="fig-err">{esc(t["render_err"])} {esc(a.get("error", ""))}</div><pre>{esc(a.get("code", ""))}</pre>')
            parts.append(f'<figure class="{"derived" if derived else "model"}">{inner}'
                         f'<figcaption><b>{esc(caption)}</b> · {esc(prov)}</figcaption></figure>')
    return (f'<!doctype html><html lang="{doc.lang}"><head><meta charset="utf-8"><title>{esc(meta["name"])} — '
            f'{esc(meta["title"])}</title><link href="{FONT_LINK}" rel="stylesheet"><style>{PDF_CSS}</style></head>'
            f'<body>{"".join(parts)}</body></html>')


def write_pdf(browser, doc: Doc, assets: dict, meta: dict, path: Path):
    try:
        import pypdf
    except ImportError:
        pypdf = None
    t = doc.t
    header = (f'<div style="font-family:Arial,sans-serif;font-size:7px;color:#{INK2};width:100%;padding:0 17mm;'
              f'display:flex;justify-content:space-between;align-items:center"><span>{logo_img("height:12px;vertical-align:middle;margin-right:6px")}{esc(meta["name"])} · {esc(meta["title"])}</span>'
              f'<span>{esc(meta["date"])}</span></div>')
    footer = (f'<div style="font-family:Arial,sans-serif;font-size:7px;color:#{INK2};width:100%;text-align:center">'
              + esc(t["footer"]).format(page='<span class="pageNumber"></span>', total='<span class="totalPages"></span>')
              + "</div>")
    opts = dict(format="A4", print_background=True, display_header_footer=True, header_template=header,
                footer_template=footer, margin={"top": "20mm", "bottom": "17mm", "left": "17mm", "right": "17mm"},
                outline=True, tagged=True)
    page = browser.new_page()

    def render(html: str, **kw) -> bytes:
        page.set_content(html, wait_until="load")
        page.evaluate("document.fonts.ready.then(() => true)")
        return page.pdf(**{**opts, **kw})

    if pypdf is None:  # single pass: cover inside the body, TOC without page numbers
        path.write_bytes(render(body_html(doc, assets, meta, None, False, True)))
        page.close()
        return
    first = render(body_html(doc, assets, meta, {}, True, False))
    pages: dict[str, str] = {}
    for i, pg in enumerate(pypdf.PdfReader(io.BytesIO(first)).pages):
        for m in re.finditer(r"zzmk(sec\d+)zz", (pg.extract_text() or "").replace(" ", "")):
            pages.setdefault(m.group(1), str(i + 1))
    body = render(body_html(doc, assets, meta, pages, False, False))
    cover_doc = (f'<!doctype html><html><head><meta charset="utf-8"><link href="{FONT_LINK}" rel="stylesheet">'
                 f'<style>{PDF_CSS}</style></head><body>{cover_html(meta, t)}</body></html>')
    cover = render(cover_doc, display_header_footer=False, margin={"top": "0", "bottom": "0", "left": "0", "right": "0"},
                   outline=False, page_ranges="1")
    page.close()
    w = pypdf.PdfWriter()
    w.append(pypdf.PdfReader(io.BytesIO(cover)), import_outline=False)
    w.append(pypdf.PdfReader(io.BytesIO(body)))
    w.add_metadata({"/Title": f'{meta["name"]} — {meta["title"]}', "/Author": "repo-analyzer",
                    "/Subject": meta["one_liner"][:200], "/Creator": "repo-analyzer (Chromium)"})
    w.page_mode = "/UseOutlines"
    with open(path, "wb") as fh:
        w.write(fh)


# ============================================================================ DOCX (python-docx)

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
             "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
             "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
             "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
             "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em",
             "lang", "eastAsianLayout", "specVanish", "oMath"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
               "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
              "tcFitText", "vAlign", "hideMark"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader",
              "tblCellSpacing", "jc", "hidden"]
SETTINGS_AFTER_UPDATEFIELDS = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                               "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                               "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
                               "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol",
                               "listSeparator"]
BODY_FONT, MONO_FONT = "Calibri", "Consolas"


class Docx:
    def __init__(self, doc: Doc, assets: dict, meta: dict):
        from docx import Document
        from docx.shared import Mm
        self.m, self.assets, self.meta, self.t = doc, assets, meta, doc.t
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(20)
        sec.top_margin, sec.bottom_margin = Mm(22), Mm(20)
        sec.header_distance = sec.footer_distance = Mm(10)
        self.text_w = sec.page_width - sec.left_margin - sec.right_margin
        self.bm = 0

    # ---- low-level xml helpers
    @staticmethod
    def el(tag: str, **attrs):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        e = OxmlElement(f"w:{tag}")
        for k, v in attrs.items():
            e.set(qn(f"w:{k}"), str(v))
        return e

    @staticmethod
    def put(parent, child, order: list[str]):
        """insert child respecting the schema order; replaces an existing element with the same tag"""
        from docx.oxml.ns import qn
        tag = child.tag.split("}")[1]
        old = parent.find(qn(f"w:{tag}"))
        if old is not None:
            parent.replace(old, child)
            return child
        later = order[order.index(tag) + 1:] if tag in order else []
        for c in parent:
            ctag = c.tag.split("}")[1] if "}" in c.tag else c.tag
            if ctag in later:
                c.addprevious(child)
                return child
        parent.append(child)
        return child

    def fonts(self, rpr, name: str):
        self.put(rpr, self.el("rFonts", ascii=name, hAnsi=name, cs=name, eastAsia=name), RPR_ORDER)

    def style_font(self, style, name: str, size: float, bold=None, color: str | None = INK, italic=None):
        from docx.shared import Pt, RGBColor
        self.fonts(style.element.get_or_add_rPr(), name)
        style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if italic is not None:
            style.font.italic = italic
        if color:
            style.font.color.rgb = RGBColor.from_string(color)

    def border(self, ppr, side: str, sz: int, color: str, val="single", space=4):
        from docx.oxml.ns import qn
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = self.put(ppr, self.el("pBdr"), PPR_ORDER)
        pbdr.append(self.el(side, val=val, sz=sz, space=space, color=color))

    # ---- styles
    def styles(self):
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt, Mm
        st = self.d.styles
        normal = st["Normal"]
        self.style_font(normal, BODY_FONT, 10)
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.15
        self.put(normal.element.get_or_add_rPr(),
                 self.el("lang", val="tr-TR" if self.m.lang == "tr" else "en-US"), RPR_ORDER)
        for name, size, color, before, after in (("Heading 1", 18, INK, 0, 10), ("Heading 2", 13, TEAL, 14, 5),
                                                 ("Heading 3", 11, INK, 10, 3)):
            s = st[name]
            self.style_font(s, BODY_FONT, size, bold=True, color=color, italic=False)
            pf = s.paragraph_format
            pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True
        st["Heading 1"].paragraph_format.page_break_before = True
        self.border(st["Heading 1"].element.get_or_add_pPr(), "bottom", 12, INK, space=4)
        self.style_font(st["Title"], BODY_FONT, 34, bold=True)
        tpp = st["Title"].element.get_or_add_pPr()
        for b in list(tpp):
            if b.tag.endswith("pBdr"):
                tpp.remove(b)
        st["Title"].paragraph_format.space_after = Pt(8)
        self.style_font(st["Caption"], BODY_FONT, 8.5, bold=False, color=INK2, italic=True)
        st["Caption"].paragraph_format.space_after = Pt(10)

        def add(name, size, color=INK, font=BODY_FONT, italic=None, bold=None, after=4, base="Normal"):
            try:
                s = st[name]
            except KeyError:
                s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = st[base]
            self.style_font(s, font, size, bold=bold, color=color, italic=italic)
            s.paragraph_format.space_after = Pt(after)
            return s

        add("Lead", 12.5, after=8)
        add("Refs", 8.5, after=6)
        add("Note", 8.8, color=INK2).paragraph_format.keep_with_next = True
        for nm, dashed in (("Prov Derived", False), ("Prov Model", True)):
            s = add(nm, 8, color=INK2, italic=True, after=8)
            s.paragraph_format.keep_with_next = True
            s.paragraph_format.left_indent = Mm(2.5)
            self.border(s.element.get_or_add_pPr(), "left", 12, RULE if not dashed else INK2,
                        val="dashed" if dashed else "single", space=6)
        code = add("Code Block", 8.2, font=MONO_FONT, after=8)
        code.paragraph_format.line_spacing = 1.0
        code.paragraph_format.left_indent = Mm(1.5)
        self.put(code.element.get_or_add_pPr(), self.el("shd", val="clear", color="auto", fill="F3F5F0"), PPR_ORDER)
        add("Table Text", 8.5, after=0).paragraph_format.line_spacing = 1.05
        add("Cover Kicker", 10, color=TEAL, bold=True, after=6)
        add("TOC Title", 18, bold=True, after=12)
        for lvl, size, bold, ind in ((1, 10.5, True, 0), (2, 9.5, False, 5)):
            s = add(f"toc {lvl}", size, bold=bold, after=2)
            s.element.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}styleId", f"TOC{lvl}")
            s.paragraph_format.left_indent = Mm(ind)
            s.paragraph_format.space_before = Pt(5 if lvl == 1 else 0)

    # ---- inline content
    def hyperlink(self, par, text: str, url: str | None = None, anchor: str | None = None, mono=True,
                  color=TEAL, size: float | None = None, strike=False):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        h = self.el("hyperlink")
        if url:
            h.set(qn("r:id"), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
        if anchor:
            h.set(qn("w:anchor"), anchor)
        r, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
        if mono:
            self.fonts(rpr, MONO_FONT)
        if strike:
            self.put(rpr, self.el("strike"), RPR_ORDER)
        if color:
            self.put(rpr, self.el("color", val=color), RPR_ORDER)
        if size:
            self.put(rpr, self.el("sz", val=int(size * 2)), RPR_ORDER)
        tx = OxmlElement("w:t")
        tx.text = text
        tx.set(qn("xml:space"), "preserve")
        r.append(rpr)
        r.append(tx)
        (h if (url or anchor) else par._p).append(r)
        if url or anchor:
            par._p.append(h)

    def runs(self, par, runs, size: float | None = None):
        from docx.shared import Pt, RGBColor
        mono = (size or 9.5) - 1.2
        for r in runs:
            if not r:
                continue
            if isinstance(r, str):
                run = par.add_run(r)
            else:
                k = r[0]
                if k == "ref":
                    _, text, ok, url = r
                    self.hyperlink(par, text, url=url if ok else None, size=mono,
                                   color=TEAL if ok else BRICK, strike=not ok)
                    continue
                if k == "sev":
                    label = r[2].replace("i", "İ").upper() if self.m.lang == "tr" else r[2].upper()
                    run = par.add_run(f" {label} ")
                    run.bold = True
                    run.font.size = Pt(7)
                    run.font.color.rgb = RGBColor.from_string("FFFFFF")
                    self.put(run._r.get_or_add_rPr(), self.el("shd", val="clear", color="auto",
                                                               fill=SEV_COLOR.get(r[1], INK2)), RPR_ORDER)
                    continue
                run = par.add_run(r[1])
                if k == "b":
                    run.bold = True
                elif k == "i":
                    run.italic = True
                elif k == "muted":
                    run.font.color.rgb = RGBColor.from_string(INK2)
                elif k == "code":
                    self.fonts(run._r.get_or_add_rPr(), MONO_FONT)
                    run.font.size = Pt(mono)
                    continue
            if size:
                run.font.size = Pt(size)

    def field(self, par, instr: str, placeholder: str = "1"):
        from docx.oxml.ns import qn
        par.add_run()._r.append(self.el("fldChar", fldCharType="begin"))
        it = self.el("instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        par.add_run()._r.append(it)
        par.add_run()._r.append(self.el("fldChar", fldCharType="separate"))
        par.add_run(placeholder)
        par.add_run()._r.append(self.el("fldChar", fldCharType="end"))

    def bookmark(self, par, name: str):
        self.bm += 1
        start = self.el("bookmarkStart", id=self.bm, name=name)
        idx = 1 if par._p.pPr is not None else 0
        par._p.insert(idx, start)
        par._p.append(self.el("bookmarkEnd", id=self.bm))

    # ---- blocks
    def table(self, headers, rows, widths, nums):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Emu
        n = len(widths)
        tbl = self.d.add_table(rows=0, cols=n)
        tbl.autofit = False
        colw = [Emu(int(self.text_w * f)) for f in widths]
        tpr = tbl._tbl.tblPr
        self.put(tpr, self.el("tblW", w=int(self.text_w / 635), type="dxa"), TBLPR_ORDER)
        b = self.el("tblBorders")
        for side, val, sz, color in (("top", "single", 8, INK), ("left", "nil", 0, "auto"), ("bottom", "single", 4, RULE),
                                     ("right", "nil", 0, "auto"), ("insideH", "single", 4, RULE), ("insideV", "nil", 0, "auto")):
            b.append(self.el(side, val=val, sz=sz, space=0, color=color))
        self.put(tpr, b, TBLPR_ORDER)
        self.put(tpr, self.el("tblLayout", type="fixed"), TBLPR_ORDER)
        mar = self.el("tblCellMar")
        for side, w in (("top", 45), ("left", 85), ("bottom", 45), ("right", 85)):
            mar.append(self.el(side, w=w, type="dxa"))
        self.put(tpr, mar, TBLPR_ORDER)
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), colw):
            gc.set(qn("w:w"), str(int(w / 635)))

        def fill(row_cells, cells, header=False):
            for j, (cell, content) in enumerate(zip(row_cells, cells)):
                cell.width = colw[j]
                par = cell.paragraphs[0]
                par.style = self.d.styles["Table Text"]
                if j in nums:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if header:
                    run = par.add_run(content)
                    run.bold = True
                    tcpr = cell._tc.get_or_add_tcPr()
                    bd = self.el("tcBorders")
                    bd.append(self.el("bottom", val="single", sz=6, space=0, color=INK))
                    self.put(tcpr, bd, TCPR_ORDER)
                    self.put(tcpr, self.el("shd", val="clear", color="auto", fill="E8ECE6"), TCPR_ORDER)
                else:
                    self.runs(par, content)

        if headers:
            row = tbl.add_row()
            trpr = row._tr.get_or_add_trPr()
            self.put(trpr, self.el("cantSplit"), TRPR_ORDER)
            self.put(trpr, self.el("tblHeader"), TRPR_ORDER)
            fill(row.cells, headers, header=True)
        for cells in rows:
            row = tbl.add_row()
            self.put(row._tr.get_or_add_trPr(), self.el("cantSplit"), TRPR_ORDER)
            fill(row.cells, cells)
        self.d.add_paragraph().paragraph_format.space_after = 0

    def figure(self, key, caption, derived, source):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Emu, Mm
        a, t = self.assets.get(key, {}), self.t
        prov = t["derived"] if derived else f'{t["model"]}{" (" + source + ")" if source else ""}'
        if a.get("png"):
            w_px, h_px = a["w"] + 28, a["h"] + 28   # stage padding
            width = min(self.text_w, Emu(int(w_px / 96 * 914400)))
            max_h = Mm(200)
            if width * h_px / w_px > max_h:
                width = Emu(int(max_h * w_px / h_px))
            self.d.add_picture(io.BytesIO(a["png"]), width=Emu(int(width)))
            pic = self.d.paragraphs[-1]
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next = True
            pic.paragraph_format.space_after = 0
        else:
            par = self.d.add_paragraph(style="Note")
            self.runs(par, [("b", t["render_err"]), " " + a.get("error", "")])
            self.d.add_paragraph(a.get("code", ""), style="Code Block")
        cap_p = self.d.add_paragraph(style="Caption")
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.runs(cap_p, [("b", caption), f" · {prov}"])

    def cover(self):
        from docx.shared import Emu, Mm, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        meta, t = self.meta, self.t
        if logo():
            self.d.add_picture(io.BytesIO(logo()), height=Mm(12))
        c = self.assets.get("__cover__", {})
        if c.get("png"):
            self.d.add_picture(io.BytesIO(c["png"]), width=Emu(int(self.text_w)))
            self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_before = Mm(10)
        self.d.add_paragraph(meta["kicker"], style="Cover Kicker")
        self.d.add_paragraph(meta["name"], style="Title")
        if meta["one_liner"]:
            self.d.add_paragraph(meta["one_liner"], style="Lead")
        rule = self.d.add_paragraph()
        self.border(rule._p.get_or_add_pPr(), "bottom", 12, INK, space=1)
        rule.paragraph_format.space_after = Pt(8)
        self.table(None, [[[("muted", a)], [b]] for a, b in meta["rows"]], [0.3, 0.7], set())
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_before = Mm(12)
        self.d.add_paragraph(t["legend_solid"], style="Prov Derived")
        self.d.add_paragraph(t["legend_dashed"], style="Prov Model")

    def toc(self):
        from docx.enum.text import WD_BREAK
        heads = [b for b in self.m.blocks if b[0] == "h" and b[1] <= 2]
        title = self.d.add_paragraph(self.t["toc"], style="TOC Title")
        title.paragraph_format.page_break_before = True
        if not heads:
            return
        from docx.oxml.ns import qn
        first = True
        par = None
        for b in heads:
            par = self.d.add_paragraph(style=f"toc {b[1]}")
            if first:
                par.add_run()._r.append(self.el("fldChar", fldCharType="begin"))
                it = self.el("instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-2" \\h \\z \\u '
                par.add_run()._r.append(it)
                par.add_run()._r.append(self.el("fldChar", fldCharType="separate"))
                first = False
            self.hyperlink(par, b[2], anchor=f"_Toc_{b[3]}", mono=False, color=INK)
        par.add_run()._r.append(self.el("fldChar", fldCharType="end"))
        settings = self.d.settings.element
        uf = self.el("updateFields", val="true")
        for c in settings:
            if c.tag.split("}")[1] in SETTINGS_AFTER_UPDATEFIELDS:
                c.addprevious(uf)
                break
        else:
            settings.append(uf)

    def header_footer(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
        for name in ("Header", "Footer"):   # template tab stops (3.25", 6.5") would pull the date to the middle
            ppr = self.d.styles[name].element.pPr
            if ppr is not None and ppr.find(qn("w:tabs")) is not None:
                ppr.remove(ppr.find(qn("w:tabs")))
        zoom = self.d.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        sec = self.d.sections[0]
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(self.text_w, WD_TAB_ALIGNMENT.RIGHT)
        if logo():
            hp.add_run().add_picture(io.BytesIO(logo()), height=Mm(4))
            hp.add_run("  ")
        for text in (f'{self.meta["name"]} · {self.meta["title"]}', "\t", self.meta["date"]):
            r = hp.add_run(text)
            r.font.size = Pt(7.5)
            r.font.color.rgb = RGBColor.from_string(INK2)
        self.border(hp._p.get_or_add_pPr(), "bottom", 4, RULE, space=4)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = self.t["footer"]
        for part in re.split(r"(\{page\}|\{total\})", fmt):
            if part == "{page}":
                self.field(fp, "PAGE")
            elif part == "{total}":
                self.field(fp, "NUMPAGES")
            elif part:
                fp.add_run(part)
        for r in fp.runs:
            r.font.size = Pt(7.5)
            r.font.color.rgb = RGBColor.from_string(INK2)

    def write(self, path: Path):
        from docx.shared import Pt
        style_of = {"lead": "Lead", "note": "Note", "refs": "Refs", "prov-derived": "Prov Derived",
                    "prov-model": "Prov Model"}
        self.styles()
        self.header_footer()
        self.cover()
        self.toc()
        for b in self.m.blocks:
            k = b[0]
            if k == "h":
                par = self.d.add_paragraph(b[2], style=f"Heading {b[1]}")
                if b[1] <= 2:
                    self.bookmark(par, f"_Toc_{b[3]}")
            elif k == "p":
                par = self.d.add_paragraph(style=style_of.get(b[2], "Normal"))
                self.runs(par, b[1])
            elif k == "bullets":
                for item in b[1]:
                    par = self.d.add_paragraph(style="List Bullet")
                    par.paragraph_format.space_after = Pt(2)
                    self.runs(par, item)
            elif k == "code":
                self.d.add_paragraph(b[1], style="Code Block")
            elif k == "table":
                self.table(b[1], b[2], b[3], b[4])
            elif k == "figure":
                self.figure(*b[1:])
        cp = self.d.core_properties
        cp.title = f'{self.meta["name"]} — {self.meta["title"]}'
        cp.subject = self.meta["one_liner"][:250]
        cp.author = "repo-analyzer"
        cp.language = "tr-TR" if self.m.lang == "tr" else "en-US"
        self.d.save(str(path))


# ============================================================================ entry points


def export(facts: dict, data: dict, t: dict, lang: str, refs: rr.Refs, dest: Path, formats, mermaid_js=None,
           verbose=False) -> list[Path]:
    from playwright.sync_api import sync_playwright
    doc = build(facts, data, t, lang, refs)
    meta = meta_of(facts, data, t, lang)
    base = dest / f'{re.sub(r"[^A-Za-z0-9._-]+", "-", facts["repo"]["name"])}-analysis-{lang}'
    figs = dict(doc.figs)
    if "docx" in formats:
        figs["__cover__"] = rr.contour_svg(facts["repo"]["name"], 1200, 380)
    written = []
    with sync_playwright() as p:
        try:
            browser = p.chromium.launch()
        except Exception as e:
            raise RuntimeError(f"Chromium not available ({str(e).splitlines()[0]}); run: python -m playwright install chromium")
        try:
            assets = render_assets(browser, figs, mermaid_js, want_png="docx" in formats)
            bad = [k for k, a in assets.items() if "error" in a]
            if bad:
                print(f"[docs] {lang}: {len(bad)} diagram(s) failed to render; their source is shown instead")
                if verbose:
                    for k in bad:
                        print("   ", assets[k]["error"])
            if "pdf" in formats:
                write_pdf(browser, doc, assets, meta, base.with_suffix(".pdf"))
                written.append(base.with_suffix(".pdf"))
        finally:
            browser.close()
    if "docx" in formats:
        Docx(doc, assets, meta).write(base.with_suffix(".docx"))
        written.append(base.with_suffix(".docx"))
    return sorted(written, key=lambda f: f.suffix != ".docx")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--langs", default=None)
    ap.add_argument("--formats", default="docx,pdf")
    ap.add_argument("--mermaid-js", default=None)
    ap.add_argument("--repo-url", default=None)
    ap.add_argument("--logo", default=None)
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args()
    args = ["--out", str(a.out), "--formats", a.formats] + (["--langs", a.langs] if a.langs else []) \
        + (["--mermaid-js", a.mermaid_js] if a.mermaid_js else []) + (["--repo-url", a.repo_url] if a.repo_url else []) \
        + (["--logo", a.logo] if a.logo else []) + (["-v"] if a.verbose else [])
    sys.argv = [str(Path(rr.__file__))] + args
    return rr.main()


if __name__ == "__main__":
    sys.exit(main())
